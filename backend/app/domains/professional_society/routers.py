import asyncio
import hashlib
import io
import json
import os
import time
import uuid
import zipfile
from datetime import datetime as _datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from beanie import PydanticObjectId
from bson import ObjectId
from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field

from ...config import get_settings
from ...core.dependencies import (
    get_current_user,
    require_role,
    require_society_access,
    require_society_event_access,
)
from ...models.attendance import Attendance
from ...models.attendance_session import AttendanceSession
from ...models.certificate import Certificate, CertSnapshot, CertStatus
from ...models.credit_rule import CreditRule
from ...models.email_log import EmailLog, EmailStatus
from ...models.event_registration import EventRegistration
from ...models.field_position import FieldPosition
from ...models.participant import Participant, ParticipantSource
from ...models.professional_society import (
    DEFAULT_SOCIETY_OFFICE_BEARER_POSITIONS,
    ProfessionalSociety,
    SocietyAssets,
)
from ...models.society_asset import SocietyAsset
from ...models.society_certificate import SocietyCertificate
from ...models.society_attendance import SocietyAttendance
from ...models.society_attendance_session import SocietyAttendanceSession
from ...models.society_participant import SocietyParticipant
from ...models.society_field_position import SocietyFieldPosition
from ...models.role_template_preset import RoleTemplatePreset
from ...models.society_event import SocietyEvent, SocietyEventAssets, SocietyEventStatus
from ...models.student_society_membership import (
    SocietyMembershipStatus,
    StudentSocietyMembership,
)
from ...models.template import Template
from ...models.user import User, UserRole
from ...routers.attendance import QR_VALIDITY_SECONDS, GenerateQRResponse, _store_nonce
from ...routers.certificates import (
    _generate_one_limited,
    _normalize_role_key,
    _resolve_event_category_role,
    _send_email_limited,
)
from ...schemas.certificate import CertificateResponse, GenerateResponse
from ...schemas.credit import CreditRuleResponse
from ...schemas.participant import (
    FieldMappingRequest,
    ParticipantCreate,
    ParticipantResponse,
    ParticipantTypeUpdate,
    UploadResponse,
)
from ...schemas.professional_society import SocietyResponse
from ...schemas.template import TemplateResponse
from ...services.cert_number import generate_cert_number
from ...services.email_service import get_daily_sent_count, send_certificate_email
from ...services.excel_service import generate_excel_template, get_active_role_names, parse_participants_excel
from ...services.png_generator import generate_certificate_from_role_preset, generate_certificate_pillow
from ...services.signature_service import process_signature, save_logo
from ...services.storage_service import (
    save_cert_png,
    storage_path_to_url,
    storage_url_to_path,
)

settings = get_settings()

router = APIRouter(prefix="/societies", tags=["Professional Societies"])
coordinator_router = APIRouter(prefix="/society-coordinator", tags=["Society Coordinator"])


def _issued_status_candidates() -> list[str]:
    return [
        CertStatus.GENERATED.value,
        CertStatus.EMAILED.value,
        "GENERATED",
        "EMAILED",
        "CertStatus.GENERATED",
        "CertStatus.EMAILED",
    ]


async def _count_event_certificates(event_id: PydanticObjectId) -> int:
    oid = ObjectId(str(event_id))
    c1 = await Certificate.find({
        "$or": [{"event_id": oid}, {"event_id": str(event_id)}],
        "status": {"$in": _issued_status_candidates()},
    }).count()
    c2 = await SocietyCertificate.find({
        "$or": [{"event_id": oid}, {"event_id": str(event_id)}],
        "status": {"$in": _issued_status_candidates()},
    }).count()
    if max(c1, c2) > 0:
        return max(c1, c2)
    tot1 = await Certificate.find({
        "$or": [{"event_id": oid}, {"event_id": str(event_id)}],
    }).count()
    tot2 = await SocietyCertificate.find({
        "$or": [{"event_id": oid}, {"event_id": str(event_id)}],
    }).count()
    return max(tot1, tot2)


def _society_response(s: ProfessionalSociety) -> SocietyResponse:
    assets = getattr(s, "assets", None)
    return SocietyResponse(
        id=str(s.id),
        name=s.name,
        slug=s.slug,
        slug_locked=getattr(s, "slug_locked", False) or False,
        contact_email=s.contact_email or "",
        is_active=s.is_active,
        created_at=s.created_at,
        logo_url=getattr(assets, "logo_url", None) if assets else None,
        signature_url=getattr(assets, "signature_url", None) if assets else None,
        signature2_url=getattr(assets, "signature2_url", None) if assets else None,
    )


def _society_assets_ready(society: ProfessionalSociety) -> bool:
    assets = getattr(society, "assets", None)
    return bool(
        assets
        and (assets.logo_path or assets.logo_url)
        and (assets.signature_path or assets.signature_url)
    )


async def _get_society_positions(society: ProfessionalSociety) -> list[str]:
    return (
        list(society.office_bearer_positions)
        if society and getattr(society, "office_bearer_positions", None)
        else list(DEFAULT_SOCIETY_OFFICE_BEARER_POSITIONS)
    )


def _fp_dict(fp: FieldPosition) -> dict:
    return {
        "id": str(fp.id),
        "event_id": str(fp.event_id),
        "cert_type": fp.cert_type,
        "template_filename": fp.template_filename,
        "column_positions": fp.column_positions,
        "asset_positions": fp.asset_positions,
        "display_width": fp.display_width,
        "confirmed": fp.confirmed,
        "created_at": fp.created_at,
    }


def _participant_resp(p: Participant) -> ParticipantResponse:
    return ParticipantResponse(
        id=str(p.id),
        event_id=str(p.event_id),
        club_id=str(p.club_id),
        email=p.email,
        registration_number=p.registration_number,
        cert_type=p.cert_type,
        fields=p.fields or {},
        field_mapping=p.field_mapping or {},
        source=p.source.value if hasattr(p.source, "value") else str(p.source),
        verified=p.verified,
        team_id=getattr(p, "team_id", None),
        team_name=getattr(p, "team_name", None),
        team_role=getattr(p, "team_role", None),
        registered_at=p.registered_at,
    )


def _cert_resp(c: Any, latest_error_by_cert: Optional[Dict[str, str]] = None) -> CertificateResponse:
    snapshot_dict: Dict[str, Any] = {}
    if hasattr(c, "snapshot") and c.snapshot:
        if hasattr(c.snapshot, "model_dump"):
            snapshot_dict = c.snapshot.model_dump()
        elif hasattr(c.snapshot, "dict"):
            snapshot_dict = c.snapshot.dict()
        elif isinstance(c.snapshot, dict):
            snapshot_dict = c.snapshot

    name = None
    email = None
    cert_type = None
    if hasattr(c, "snapshot") and c.snapshot:
        name = getattr(c.snapshot, "name", None) or snapshot_dict.get("name")
        email = getattr(c.snapshot, "email", None) or snapshot_dict.get("email")
        cert_type = getattr(c.snapshot, "cert_type", None) or snapshot_dict.get("cert_type")

    status_str = c.status.value if hasattr(c.status, "value") else str(c.status)
    err = (latest_error_by_cert or {}).get(str(c.id))
    file_url = storage_path_to_url(c.png_url) if getattr(c, "png_url", None) else None

    return CertificateResponse(
        id=str(c.id),
        cert_number=c.cert_number or "",
        participant_id=str(c.participant_id),
        event_id=str(c.event_id),
        template_id=str(getattr(c, "template_id", c.id) or c.id),
        club_id=str(getattr(c, "club_id", getattr(c, "society_id", ""))),
        participant_name=name,
        participant_email=email,
        cert_type=cert_type,
        snapshot=snapshot_dict,
        status=status_str,
        failure_reason=err,
        pdf_url=file_url,
        png_url=file_url,
        generated_at=getattr(c, "issued_at", None),
        issued_at=getattr(c, "issued_at", None),
        revoked_at=getattr(c, "revoked_at", None),
    )



def _extract_asset_val(obj: Any, attr: str) -> Optional[str]:
    if obj is None:
        return None
    if isinstance(obj, dict):
        return obj.get(attr)
    return getattr(obj, attr, None)


async def _get_effective_society_assets(society_id: PydanticObjectId, event_assets: Optional[Any] = None) -> dict:
    soc_asset = await SocietyAsset.find_one(SocietyAsset.society_id == society_id)
    society = await ProfessionalSociety.get(society_id)
    fallback_assets = soc_asset if soc_asset else (society.assets if society else None)

    def _val(attr: str):
        return _extract_asset_val(event_assets, attr) or _extract_asset_val(fallback_assets, attr)

    logo_url = _val("logo_url")
    logo_path = _val("logo_path")
    logo_hash = _val("logo_hash")
    signature_url = _val("signature_url")
    signature_path = _val("signature_path")
    signature_hash = _val("signature_hash")
    signature2_url = _val("signature2_url")
    signature2_path = _val("signature2_path")
    signature2_hash = _val("signature2_hash")

    return {
        "logo_url": logo_url,
        "logo_path": logo_path,
        "logo_hash": logo_hash,
        "signature_url": signature_url,
        "signature_path": signature_path,
        "signature_hash": signature_hash,
        "signature2_url": signature2_url,
        "signature2_path": signature2_path,
        "signature2_hash": signature2_hash,
        "is_configured": bool(logo_url and signature_url),
    }

def _event_resp(
    e: SocietyEvent,
    cert_count: int = 0,
    volunteers_registered: int = 0,
    team_count: int = 0,
    effective_assets: Optional[dict] = None,
) -> dict:
    return {
        "id": str(e.id),
        "society_id": str(e.society_id),
        "name": e.name,
        "description": e.description,
        "event_date": e.event_date,
        "event_time": e.event_time,
        "session": e.session,
        "slot": e.slot,
        "venue": e.venue,
        "category": e.category,
        "academic_year": e.academic_year,
        "academic_years": e.academic_years or [],
        "status": e.status.value if hasattr(e.status, "value") else str(e.status),
        "poster_url": e.poster_url,
        "is_published": e.is_published,
        "max_participants": e.max_participants,
        "participant_count": e.participant_count or 0,
        "volunteers_required": e.volunteers_required or 0,
        "volunteers_registered": volunteers_registered,
        "is_team_event": getattr(e, "is_team_event", False) or False,
        "min_team_size": getattr(e, "min_team_size", 2) or 2,
        "max_team_size": getattr(e, "max_team_size", 4) or 4,
        "max_teams": getattr(e, "max_teams", None),
        "team_count": team_count,
        "qr_generations_count": getattr(e, "qr_generations_count", 0) or 0,
        "registration_stopped": getattr(e, "registration_stopped", False) or False,
        "issued_certificates_count": cert_count,
        "template_filename": e.template_filename,
        "template_map": e.template_map or {},
        "mapping_confirmed": getattr(e, "mapping_confirmed", False) or False,
        "assets": effective_assets if effective_assets is not None else getattr(e, "assets", SocietyEventAssets()).model_dump(),
        "created_at": e.created_at,
    }


# ── Pydantic Request Models ──────────────────────────────────────────────────

class SocietyEventCreate(BaseModel):
    name: str
    description: str = Field(min_length=1)
    event_date: Optional[_datetime] = None
    event_time: Optional[str] = None
    session: Optional[str] = None
    slot: Optional[str] = None
    venue: Optional[str] = None
    category: Optional[str] = None
    academic_year: Optional[str] = None
    academic_years: List[str] = Field(default_factory=list)
    max_participants: Optional[int] = None
    volunteers_required: int = 0
    is_team_event: bool = False
    min_team_size: Optional[int] = 2
    max_team_size: Optional[int] = 4
    max_teams: Optional[int] = None
    is_published: bool = False


class SocietyEventUpdate(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    event_date: Optional[_datetime] = None
    event_time: Optional[str] = None
    session: Optional[str] = None
    slot: Optional[str] = None
    venue: Optional[str] = None
    category: Optional[str] = None
    academic_year: Optional[str] = None
    academic_years: Optional[List[str]] = None
    max_participants: Optional[int] = None
    volunteers_required: Optional[int] = None
    is_team_event: Optional[bool] = None
    min_team_size: Optional[int] = None
    max_team_size: Optional[int] = None
    max_teams: Optional[int] = None
    registration_stopped: Optional[bool] = None
    status: Optional[str] = None
    is_published: Optional[bool] = None


class ReviewMembershipRequest(BaseModel):
    action: Optional[str] = None
    status: Optional[str] = None
    review_note: Optional[str] = None


class BulkApproveMembershipsRequest(BaseModel):
    membership_ids: List[str]
    review_note: Optional[str] = None


class AssignOfficeBearerRequest(BaseModel):
    position: str
    student_id: Optional[str] = None
    student_email: Optional[str] = None


class AddPositionBody(BaseModel):
    position: str


class AllocateOfficeBearerBody(BaseModel):
    position: str
    student_id: Optional[str] = None
    student_email: Optional[str] = None


class VolunteerCountUpdate(BaseModel):
    volunteers_required: int


class VolunteerStatusUpdate(BaseModel):
    status: Optional[str] = None
    action: Optional[str] = None


class VolunteerRequestResponse(BaseModel):
    id: str
    participant_id: Optional[str] = None
    registration_id: Optional[str] = None
    student_name: str
    student_email: str
    registration_number: Optional[str] = None
    department: Optional[str] = None
    status: str
    verified: bool = False
    registered_at: Optional[_datetime] = None


class FieldPositionRequest(BaseModel):
    cert_type: str
    template_filename: Optional[str] = None
    column_positions: Dict[str, Any] = Field(default_factory=dict)
    asset_positions: Optional[Dict[str, Any]] = None
    display_width: Optional[float] = None
    confirmed: bool = False


class AssignPresetRequest(BaseModel):
    preset_id: str


class UpdatePresetSlotsRequest(BaseModel):
    slots: Dict[str, Any]


class TemplateMapRequest(BaseModel):
    template_map: Dict[str, Optional[str]]


# ══════════════════════════════════════════════════════════════════════════════
# Society Basic & Event Endpoints
# ══════════════════════════════════════════════════════════════════════════════

@router.get("", response_model=List[SocietyResponse])
async def list_active_societies(_user: User = Depends(get_current_user)):
    societies = await ProfessionalSociety.find(ProfessionalSociety.is_active == True).to_list()
    return [_society_response(s) for s in societies]


@router.get("/{society_id}", response_model=SocietyResponse)
async def get_society(society_id: PydanticObjectId, _user: User = Depends(get_current_user)):
    society = await ProfessionalSociety.get(society_id)
    if not society or not society.is_active:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Professional Society not found or inactive")
    return _society_response(society)


@router.get("/{society_id}/dashboard")
async def society_dashboard(
    society_id: PydanticObjectId,
    _user: User = Depends(require_society_access),
):
    society = await ProfessionalSociety.get(society_id)
    if not society:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Professional Society not found")

    all_events = await SocietyEvent.find(SocietyEvent.society_id == society_id).to_list()
    event_ids = [ev.id for ev in all_events]
    certs = await Certificate.find({"event_id": {"$in": event_ids}}).to_list() if event_ids else []
    certs_by_event: Dict[str, List[Certificate]] = {}
    for cert in certs:
        certs_by_event.setdefault(str(cert.event_id), []).append(cert)

    completed_events = []
    for ev in all_events:
        event_certs = certs_by_event.get(str(ev.id), [])
        if not event_certs:
            continue
        if any(cert.status != CertStatus.EMAILED for cert in event_certs):
            continue
        completed_events.append(ev)

    completed_event_ids = [ev.id for ev in completed_events]
    now = _datetime.utcnow()
    if now.month >= 7:
        sem_start = _datetime(now.year, 7, 1)
        sem_end = _datetime(now.year, 12, 31, 23, 59, 59)
    else:
        sem_start = _datetime(now.year, 1, 1)
        sem_end = _datetime(now.year, 6, 30, 23, 59, 59)

    events_this_semester = [
        ev for ev in completed_events
        if ev.event_date and sem_start <= ev.event_date <= sem_end
    ]

    total_events_this_semester = len(events_this_semester)
    cumulative_completed_events = len(completed_events)
    total_certificates_issued = sum(len(certs_by_event.get(str(ev.id), [])) for ev in completed_events)
    total_participants = await Participant.find(
        {"event_id": {"$in": completed_event_ids}},
    ).count() if completed_event_ids else 0

    recent_events_docs = sorted(completed_events, key=lambda e: e.created_at, reverse=True)[:5]
    recent_events = []
    for ev in recent_events_docs:
        p_count = await Participant.find(Participant.event_id == ev.id).count()
        c_count = len(certs_by_event.get(str(ev.id), []))
        recent_events.append({
            "event_id": str(ev.id),
            "name": ev.name,
            "event_date": ev.event_date,
            "status": ev.status.value,
            "participant_count": p_count,
            "certificates_issued_count": c_count,
        })

    active_members_count = await StudentSocietyMembership.find({
        "society_id": society_id,
        "status": SocietyMembershipStatus.APPROVED.value,
    }).count()

    society_payload = _society_response(society).model_dump()
    if not society_payload.get("contact_email"):
        coordinator = await User.find_one(
            User.society_id == society_id,
            User.role == UserRole.PROFESSIONAL_SOCIETY_COORD,
        )
        if coordinator and coordinator.email:
            society_payload["contact_email"] = coordinator.email

    return {
        "society": society_payload,
        "club": society_payload,
        "stats": {
            "total_events": total_events_this_semester,
            "events_this_semester": total_events_this_semester,
            "cumulative_events": cumulative_completed_events,
            "total_certificates_issued": total_certificates_issued,
            "total_participants": total_participants,
        },
        "society_id": str(society.id),
        "name": society.name,
        "slug": society.slug,
        "contact_email": society.contact_email or "",
        "total_events_this_semester": total_events_this_semester,
        "cumulative_completed_events": cumulative_completed_events,
        "total_certificates_issued": total_certificates_issued,
        "total_participants": total_participants,
        "recent_events": recent_events,
        "total_events": len(all_events),
        "published_events_count": len([e for e in all_events if e.is_published]),
        "active_members_count": active_members_count,
        "assets_configured": _society_assets_ready(society),
        "office_bearer_positions": await _get_society_positions(society),
    }


@router.get("/{society_id}/assets")
async def get_society_assets(
    society_id: PydanticObjectId,
    _user: User = Depends(require_society_access),
):
    society = await ProfessionalSociety.get(society_id)
    if not society:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Professional Society not found")

    soc_asset = await SocietyAsset.find_one(SocietyAsset.society_id == society_id)
    assets = soc_asset if soc_asset else society.assets
    logo_u = getattr(assets, "logo_url", None)
    sig_u = getattr(assets, "signature_url", None)
    return {
        "logo_url": logo_u,
        "logo_hash": getattr(assets, "logo_hash", None),
        "signature_url": sig_u,
        "signature_hash": getattr(assets, "signature_hash", None),
        "is_configured": bool(logo_u and sig_u),
    }


@router.post("/{society_id}/assets")
async def update_society_assets(
    society_id: PydanticObjectId,
    logo: UploadFile | None = File(None),
    signature: UploadFile | None = File(None),
    _user: User = Depends(require_society_access),
):
    if not logo and not signature:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Upload logo or signature")

    society = await ProfessionalSociety.get(society_id)
    if not society:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Professional Society not found")

    soc_asset = await SocietyAsset.find_one(SocietyAsset.society_id == society_id)
    if not soc_asset:
        soc_asset = SocietyAsset(society_id=society_id, society_slug=society.slug)

    assets = society.assets
    slug = society.slug or "society"

    if logo:
        logo_bytes = await logo.read()
        h = hashlib.md5(logo_bytes).hexdigest()
        p = save_logo(logo_bytes, slug)
        u = storage_path_to_url(p)
        assets.logo_hash = h
        assets.logo_path = p
        assets.logo_url = u
        soc_asset.logo_hash = h
        soc_asset.logo_path = p
        soc_asset.logo_url = u

    if signature:
        sig_bytes = await signature.read()
        h = hashlib.md5(sig_bytes).hexdigest()
        p = process_signature(sig_bytes, slug)
        u = storage_path_to_url(p)
        assets.signature_hash = h
        assets.signature_path = p
        assets.signature_url = u
        soc_asset.signature_hash = h
        soc_asset.signature_path = p
        soc_asset.signature_url = u

    soc_asset.updated_at = _datetime.utcnow()
    await soc_asset.save()
    await society.set({"assets": assets.model_dump()})

    return {
        "message": "Society assets updated",
        "assets": {
            "logo_url": soc_asset.logo_url,
            "signature_url": soc_asset.signature_url,
            "is_configured": bool(soc_asset.logo_url and soc_asset.signature_url),
        },
    }


@router.get("/{society_id}/credit-rules", response_model=List[CreditRuleResponse])
async def get_society_credit_rules(
    society_id: PydanticObjectId,
    _user: User = Depends(require_society_access),
):
    society = await ProfessionalSociety.get(society_id)
    if not society or not society.is_active:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Professional Society not found")
    rules = await CreditRule.find_all().to_list()
    return [
        CreditRuleResponse(
            id=str(r.id),
            cert_type=r.cert_type,
            points=r.points,
            updated_at=r.updated_at,
            updated_by=str(r.updated_by) if r.updated_by else None,
        )
        for r in rules
    ]


@router.get("/{society_id}/members")
async def get_society_members_list(
    society_id: PydanticObjectId,
    _user: User = Depends(require_society_access),
):
    members = await StudentSocietyMembership.find({
        "society_id": society_id,
        "status": SocietyMembershipStatus.APPROVED.value,
    }).to_list()
    return [
        {
            "id": str(m.id),
            "student_id": str(m.student_id),
            "student_name": m.student_name,
            "student_email": m.student_email,
            "registration_number": getattr(m, "registration_number", None),
            "office_bearer_role": m.office_bearer_role,
            "applied_at": m.applied_at,
        }
        for m in members
    ]


# ─── Events CRUD ─────────────────────────────────────────────────────────────

@router.get("/{society_id}/events")
async def get_society_events(
    society_id: PydanticObjectId,
    _user: User = Depends(require_society_access),
):
    events = await SocietyEvent.find(SocietyEvent.society_id == society_id).sort("-created_at").to_list()
    event_ids = [e.id for e in events]

    participant_regs = await EventRegistration.find(
        {"event_id": {"$in": event_ids}, "registration_type": {"$ne": "volunteer"}}
    ).to_list() if event_ids else []

    vol_regs = await EventRegistration.find(
        {"event_id": {"$in": event_ids}, "registration_type": "volunteer", "$or": [{"status": "accepted"}, {"status": "pending"}]}
    ).to_list() if event_ids else []

    reg_counts = {}
    team_counts = {}
    for pr in participant_regs:
        eid = str(pr.event_id)
        reg_counts[eid] = reg_counts.get(eid, 0) + 1
        if getattr(pr, "team_id", None) and getattr(pr, "team_role", None) == "leader":
            team_counts[eid] = team_counts.get(eid, 0) + 1

    vol_counts = {}
    for vr in vol_regs:
        eid = str(vr.event_id)
        vol_counts[eid] = vol_counts.get(eid, 0) + 1

    effective_assets = await _get_effective_society_assets(society_id, None)
    result = []
    for e in events:
        eid = str(e.id)
        actual_count = max(e.participant_count or 0, reg_counts.get(eid, 0))
        if actual_count != e.participant_count:
            await e.set({"participant_count": actual_count})
            e.participant_count = actual_count
        cert_count = await _count_event_certificates(e.id)
        evt_assets = getattr(e, "assets", None)
        assets_for_event = (
            await _get_effective_society_assets(society_id, evt_assets)
            if (evt_assets and (evt_assets.logo_url or evt_assets.signature_url))
            else effective_assets
        )
        result.append(_event_resp(
            e,
            cert_count=cert_count,
            volunteers_registered=vol_counts.get(eid, 0),
            team_count=team_counts.get(eid, 0),
            effective_assets=assets_for_event,
        ))
    return result


@router.post("/{society_id}/events", status_code=201)
async def create_society_event(
    society_id: PydanticObjectId,
    body: SocietyEventCreate,
    _user: User = Depends(require_society_access),
):
    society = await ProfessionalSociety.get(society_id)
    if not society:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Professional Society not found")

    status_val = SocietyEventStatus.ACTIVE if body.is_published else SocietyEventStatus.DRAFT
    new_event = SocietyEvent(
        society_id=society_id,
        name=body.name,
        description=body.description,
        event_date=body.event_date,
        event_time=body.event_time,
        session=body.session,
        slot=body.slot,
        venue=body.venue,
        category=body.category,
        academic_year=body.academic_year,
        academic_years=body.academic_years or [],
        status=status_val,
        max_participants=body.max_participants,
        volunteers_required=body.volunteers_required,
        is_team_event=body.is_team_event,
        min_team_size=body.min_team_size or 2,
        max_team_size=body.max_team_size or 4,
        max_teams=body.max_teams,
        is_published=body.is_published,
    )
    await new_event.insert()
    return {"message": "Event created successfully", "id": str(new_event.id), "event_id": str(new_event.id)}


@router.get("/{society_id}/events/{event_id}")
async def get_society_event_detail(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    reg_count = await EventRegistration.find(
        {"event_id": event_id, "registration_type": {"$ne": "volunteer"}}
    ).count()
    uploaded_count = await Participant.find(Participant.event_id == event_id).count()
    actual_count = max(reg_count, uploaded_count, event.participant_count or 0)
    if actual_count != event.participant_count:
        await event.set({"participant_count": actual_count})
        event.participant_count = actual_count

    mapping_confirmed = (
        await FieldPosition.find(
            FieldPosition.event_id == event_id,
            FieldPosition.confirmed == True,
        ).count()
    ) > 0
    if mapping_confirmed != getattr(event, "mapping_confirmed", False):
        await event.set({"mapping_confirmed": mapping_confirmed})
        event.mapping_confirmed = mapping_confirmed

    team_count = await EventRegistration.find(
        EventRegistration.event_id == event_id,
        EventRegistration.team_role == "leader",
    ).count()

    vol_count = await EventRegistration.find(
        EventRegistration.event_id == event_id,
        EventRegistration.registration_type == "volunteer",
        {"$or": [{"status": "accepted"}, {"status": "pending"}]},
    ).count()

    cert_count = await _count_event_certificates(event.id)
    effective_assets = await _get_effective_society_assets(society_id, getattr(event, "assets", None))
    return _event_resp(event, cert_count=cert_count, volunteers_registered=vol_count, team_count=team_count, effective_assets=effective_assets)


@router.put("/{society_id}/events/{event_id}")
@router.patch("/{society_id}/events/{event_id}")
async def update_society_event(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    body: SocietyEventUpdate,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    updates = body.model_dump(exclude_unset=True)
    if "is_published" in updates and "status" not in updates:
        updates["status"] = SocietyEventStatus.ACTIVE if updates["is_published"] else SocietyEventStatus.DRAFT

    if "status" in updates:
        try:
            status_enum = SocietyEventStatus(str(updates["status"]).lower())
            updates["status"] = status_enum
            if status_enum == SocietyEventStatus.ACTIVE:
                updates["is_published"] = True
            elif status_enum == SocietyEventStatus.DRAFT:
                updates["is_published"] = False
        except Exception:
            pass

    if updates:
        await event.set(updates)
        event = await SocietyEvent.get(event_id)

    cert_count = await _count_event_certificates(event.id)
    effective_assets = await _get_effective_society_assets(society_id, getattr(event, "assets", None))
    return {"message": "Event updated successfully", "id": str(event.id), "event": _event_resp(event, cert_count=cert_count, effective_assets=effective_assets)}


@router.delete("/{society_id}/events/{event_id}", status_code=204)
async def delete_society_event(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    await Participant.find(Participant.event_id == event_id).delete()
    await Certificate.find({
        "event_id": event_id,
        "status": {"$ne": CertStatus.EMAILED.value},
    }).delete()
    await FieldPosition.find(FieldPosition.event_id == event_id).delete()
    await event.delete()
    return None


@router.post("/{society_id}/events/{event_id}/poster")
async def upload_society_event_poster(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    poster: UploadFile = File(...),
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    content = await poster.read()
    filename = poster.filename or "poster.jpg"
    ext = Path(filename).suffix or ".jpg"
    poster_dir = settings.storage_root / "posters" / f"society_{str(society_id)}"
    poster_dir.mkdir(parents=True, exist_ok=True)
    poster_path = poster_dir / f"{str(event_id)}_poster{ext}"
    poster_path.write_bytes(content)
    poster_url = storage_path_to_url(str(poster_path))

    await event.set({
        "poster_path": str(poster_path),
        "poster_url": poster_url,
    })

    return {"message": "Poster uploaded", "poster_url": poster_url}


@router.post("/{society_id}/events/{event_id}/stop-registration")
async def stop_society_registration(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")
    await event.set({"registration_stopped": True})
    return {"message": "Registration stopped for this event"}


@router.post("/{society_id}/events/{event_id}/resume-registration")
async def resume_society_registration(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")
    await event.set({"registration_stopped": False})
    return {"message": "Registration resumed for this event"}


# ─── Volunteer Management & Excel Template ───────────────────────────────────

@router.get("/{society_id}/events/{event_id}/volunteers", response_model=List[VolunteerRequestResponse])
@router.get("/{society_id}/events/{event_id}/volunteer-requests", response_model=List[VolunteerRequestResponse])
async def list_society_volunteer_requests(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    participants = await Participant.find(
        Participant.event_id == event_id,
        Participant.cert_type == "volunteer",
    ).to_list()

    registrations = await EventRegistration.find(
        EventRegistration.event_id == event_id,
        EventRegistration.registration_type == "volunteer",
    ).to_list()
    reg_map = {r.student_email.lower(): r for r in registrations if r.student_email}

    results = []
    seen_emails = set()

    for p in participants:
        email = (p.email or "").lower()
        seen_emails.add(email)
        reg = reg_map.get(email)
        status_val = getattr(p, "status", None)
        if not status_val:
            status_val = "accepted" if p.verified else (getattr(reg, "status", "pending") if reg else "pending")

        results.append(VolunteerRequestResponse(
            id=str(p.id),
            participant_id=str(p.id),
            registration_id=str(reg.id) if reg else None,
            student_name=p.fields.get("Name") or (reg.student_name if reg else "") or email,
            student_email=p.email,
            registration_number=p.registration_number or (reg.registration_number if reg else None),
            department=p.fields.get("Department") or (reg.department if reg else None),
            status=status_val,
            verified=p.verified,
            registered_at=p.registered_at,
        ))

    for email, reg in reg_map.items():
        if email not in seen_emails:
            results.append(VolunteerRequestResponse(
                id=str(reg.id),
                participant_id=None,
                registration_id=str(reg.id),
                student_name=reg.student_name or email,
                student_email=reg.student_email,
                registration_number=reg.registration_number,
                department=reg.department,
                status=getattr(reg, "status", "pending") or "pending",
                verified=False,
                registered_at=reg.registered_at if hasattr(reg, "registered_at") else None,
            ))

    return results


@router.patch("/{society_id}/events/{event_id}/volunteers/{item_id}")
@router.patch("/{society_id}/events/{event_id}/volunteer-requests/{item_id}")
async def update_society_volunteer_status(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    item_id: str,
    body: VolunteerStatusUpdate,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    target_status = (body.status or body.action or "accepted").lower()
    if target_status not in ["accepted", "rejected", "approved"]:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Status must be 'accepted' or 'rejected'")
    if target_status == "approved":
        target_status = "accepted"

    is_verified = target_status == "accepted"

    p = None
    reg = None
    try:
        oid = ObjectId(item_id)
        p = await Participant.get(oid)
    except Exception:
        pass

    try:
        oid = ObjectId(item_id)
        reg = await EventRegistration.get(oid)
    except Exception:
        pass

    email = None
    if p:
        email = (p.email or "").lower()
        await p.set({"verified": is_verified, "status": target_status})
    if reg:
        email = email or (reg.student_email or "").lower()
        await reg.set({"status": target_status})

    if email and not reg:
        matching_reg = await EventRegistration.find_one(
            EventRegistration.event_id == event_id,
            EventRegistration.student_email == email,
        )
        if matching_reg:
            await matching_reg.set({"status": target_status})

    if email and not p:
        matching_p = await Participant.find_one(
            Participant.event_id == event_id,
            Participant.email == email,
        )
        if matching_p:
            await matching_p.set({"verified": is_verified, "status": target_status})
        elif is_verified and reg:
            new_p = Participant(
                event_id=event_id,
                club_id=society_id,
                email=email,
                registration_number=reg.registration_number,
                cert_type="volunteer",
                fields={
                    "Name": reg.student_name or "",
                    "Email": email,
                    "Registration Number": reg.registration_number or "",
                    "Department": reg.department or "",
                },
                source=ParticipantSource.REGISTRATION,
                verified=True,
                status="accepted",
            )
            await new_p.insert()

    return {"message": f"Volunteer request marked as {target_status}"}


@router.patch("/{society_id}/events/{event_id}/volunteers-count")
@router.patch("/{society_id}/events/{event_id}/volunteers-required")
async def update_society_volunteers_count(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    body: VolunteerCountUpdate,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")
    await event.set({"volunteers_required": body.volunteers_required})
    event = await SocietyEvent.get(event_id)
    cert_count = await _count_event_certificates(event.id)
    return _event_resp(event, cert_count=cert_count)


@router.get("/{society_id}/events/{event_id}/credit-rules", response_model=List[CreditRuleResponse])
async def get_society_event_credit_rules(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    rules = await CreditRule.find_all().to_list()
    return [
        CreditRuleResponse(
            id=str(r.id),
            cert_type=r.cert_type,
            points=r.points,
            updated_at=r.updated_at,
            updated_by=str(r.updated_by) if r.updated_by else None,
        )
        for r in rules
    ]


@router.get("/{society_id}/events/{event_id}/excel-template")
async def download_society_excel_template(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    roles = await get_active_role_names()
    buf = generate_excel_template(roles=roles)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=certificate_template.xlsx"},
    )


# ─── Field Positions & Templates ─────────────────────────────────────────────

@router.post("/{society_id}/events/{event_id}/field-positions")
async def save_society_field_positions(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    body: FieldPositionRequest,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    fp = await FieldPosition.find_one(
        FieldPosition.event_id == event_id,
        FieldPosition.cert_type == body.cert_type,
    )
    if fp:
        fp.template_filename = body.template_filename
        fp.column_positions = body.column_positions
        if body.asset_positions is not None:
            fp.asset_positions = body.asset_positions
        if body.display_width is not None:
            fp.display_width = body.display_width
        fp.confirmed = body.confirmed
        await fp.save()
    else:
        fp = FieldPosition(
            event_id=event_id,
            cert_type=body.cert_type,
            template_filename=body.template_filename,
            column_positions=body.column_positions,
            asset_positions=body.asset_positions or {},
            display_width=body.display_width,
            confirmed=body.confirmed,
        )
        await fp.insert()

    # Update event template_map and confirmation
    current_map = dict(event.template_map or {})
    if body.template_filename:
        current_map[body.cert_type] = body.template_filename
        updates = {"template_map": current_map}
        if body.confirmed:
            updates["mapping_confirmed"] = True
        await event.set(updates)

    return _fp_dict(fp)


@router.get("/{society_id}/events/{event_id}/field-positions")
async def get_society_all_field_positions(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    docs = await FieldPosition.find(FieldPosition.event_id == event_id).to_list()
    return [_fp_dict(fp) for fp in docs]


@router.get("/{society_id}/events/{event_id}/field-positions/{cert_type}")
async def get_society_field_positions_for_cert_type(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    cert_type: str,
    _user: User = Depends(require_society_event_access),
):
    fp = await FieldPosition.find_one(
        FieldPosition.event_id == event_id,
        FieldPosition.cert_type == cert_type,
    )
    if not fp:
        raise HTTPException(status.HTTP_404_NOT_FOUND, f"No field positions for cert_type '{cert_type}'")
    return _fp_dict(fp)


@router.post("/{society_id}/events/{event_id}/template-map")
async def update_society_template_map(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    body: TemplateMapRequest,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")
    await event.set({"template_map": body.template_map})
    return {"message": "Template map updated", "template_map": body.template_map}


@router.post("/{society_id}/events/{event_id}/template-confirm")
async def confirm_society_template_mapping(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")
    await event.set({"mapping_confirmed": True})
    return {"message": "Template mapping confirmed"}


@router.get("/{society_id}/templates", response_model=List[TemplateResponse])
async def list_society_templates(
    society_id: PydanticObjectId,
    _user: User = Depends(require_society_access),
):
    templates = await Template.find_all().to_list()
    return [
        TemplateResponse(
            id=str(t.id),
            name=t.name,
            cert_type=t.cert_type,
            fields=t.fields,
            created_at=t.created_at,
        )
        for t in templates
    ]


@router.get("/{society_id}/templates/presets", response_model=List[TemplateResponse])
@router.get("/{society_id}/events/{event_id}/templates/presets", response_model=List[TemplateResponse])
async def list_society_preset_templates(
    society_id: PydanticObjectId,
    event_id: Optional[PydanticObjectId] = None,
    _user: User = Depends(require_society_access),
):
    presets = await RoleTemplatePreset.find(RoleTemplatePreset.is_active == True).to_list()
    return [
        TemplateResponse(
            id=str(p.id),
            name=p.display_name or p.role_name,
            cert_type=p.role_name,
            fields=[],
            created_at=_datetime.utcnow(),
        )
        for p in presets
    ]


@router.post("/{society_id}/events/{event_id}/templates/assign-preset")
async def assign_society_preset(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    body: AssignPresetRequest,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    preset = await RoleTemplatePreset.get(ObjectId(body.preset_id))
    if not preset:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Preset not found")

    current_map = dict(event.template_map or {})
    current_map[preset.role_name] = preset.template_filename
    await event.set({
        "template_map": current_map,
        "mapping_confirmed": True,
    })
    return {"message": f"Preset '{preset.role_name}' assigned successfully"}


@router.post("/{society_id}/events/{event_id}/templates/update-preset-slots")
@router.patch("/{society_id}/events/{event_id}/templates/update-preset-slots")
async def update_society_preset_slots(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    body: UpdatePresetSlotsRequest,
    _user: User = Depends(require_society_event_access),
):
    return {"message": "Preset slots updated successfully"}


# ─── Participants Management ─────────────────────────────────────────────────

@router.get("/{society_id}/events/{event_id}/participants", response_model=List[ParticipantResponse])
async def list_society_participants(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    source: Optional[str] = None,
    verified: Optional[bool] = None,
    cert_type: Optional[str] = None,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    # Auto-sync registered participants from EventRegistration into Participant collection
    portal_regs = await EventRegistration.find(
        {"event_id": event_id, "registration_type": {"$ne": "volunteer"}}
    ).to_list()

    # Query attendance records to accurately reflect who has actually attended
    attendances = await Attendance.find(Attendance.event_id == event_id).to_list()
    attended_emails = {
        (a.student_email or "").strip().lower()
        for a in attendances
        if a.student_email
    }

    if portal_regs:
        existing_parts = await Participant.find(Participant.event_id == event_id).to_list()
        existing_emails = {p.email.lower() for p in existing_parts if p.email}

        for reg in portal_regs:
            email = (reg.student_email or "").strip().lower()
            if not email or email in existing_emails:
                continue
            default_role = _resolve_event_category_role("participant", getattr(event, "category", None))
            has_attended = email in attended_emails
            new_p = Participant(
                event_id=event_id,
                club_id=society_id,
                email=email,
                registration_number=reg.registration_number,
                cert_type=default_role,
                fields={
                    "Name": reg.student_name or "",
                    "Email": email,
                    "Registration Number": reg.registration_number or "",
                    "Department": reg.department or "",
                    "Role": default_role.replace("_", " ").title(),
                },
                source=ParticipantSource.REGISTRATION,
                verified=has_attended,
                status="accepted" if has_attended else "registered",
                team_id=getattr(reg, "team_id", None),
                team_name=getattr(reg, "team_name", None),
                team_role=getattr(reg, "team_role", None),
            )
            await new_p.insert()
            existing_emails.add(email)

        # For existing participants from registration, sync verified flag if not manually overridden
        for p in existing_parts:
            if p.source == ParticipantSource.REGISTRATION and not getattr(p, "manually_verified", False):
                p_email = (p.email or "").strip().lower()
                actual_attended = p_email in attended_emails
                if p.verified != actual_attended:
                    await p.set({"verified": actual_attended, "status": "accepted" if actual_attended else "registered"})

    query: Dict[str, Any] = {"event_id": event_id}
    if source:
        query["source"] = source
    if verified is not None:
        query["verified"] = verified
    if cert_type:
        query["cert_type"] = cert_type

    participants = await Participant.find(query).to_list()
    return [_participant_resp(p) for p in participants]


@router.post("/{society_id}/events/{event_id}/participants", response_model=ParticipantResponse, status_code=201)
async def create_society_participant(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    body: ParticipantCreate,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    email = body.email.strip().lower()
    existing = await Participant.find_one(
        Participant.event_id == event_id,
        Participant.email == email,
    )
    if existing:
        raise HTTPException(status.HTTP_409_CONFLICT, "Participant already registered")

    p = Participant(
        event_id=event_id,
        club_id=society_id,
        email=email,
        registration_number=body.registration_number,
        cert_type=body.cert_type,
        fields={**body.fields, **({"Name": body.name} if body.name else {})},
        source=ParticipantSource.MANUAL,
        verified=True,
    )
    await p.insert()

    new_count = (event.participant_count or 0) + 1
    await event.set({"participant_count": new_count})

    return _participant_resp(p)


@router.get("/{society_id}/events/{event_id}/participants/columns")
async def get_society_participant_columns(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    participants = await Participant.find(Participant.event_id == event_id).to_list()
    columns: list[str] = []
    sample: dict[str, str] = {}
    for p in participants:
        fields = p.fields or {}
        if not sample and fields:
            sample = fields
        for col in fields:
            if col not in columns:
                columns.append(col)
    return {"columns": columns, "sample": sample}


@router.get("/{society_id}/events/{event_id}/participants/mapping")
async def get_society_participant_field_mapping(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    participant = await Participant.find_one({
        "event_id": event_id,
        "field_mapping": {"$ne": {}},
    })
    return {"mapping": participant.field_mapping if participant else {}}


@router.post("/{society_id}/events/{event_id}/participants/mapping")
async def apply_society_participant_field_mapping(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    body: FieldMappingRequest,
    _user: User = Depends(require_society_event_access),
):
    participants = await Participant.find(Participant.event_id == event_id).to_list()
    count = 0
    for p in participants:
        await p.set({"field_mapping": body.mapping})
        count += 1
    return {"message": f"Mapping applied to {count} participant(s)"}


@router.post("/{society_id}/events/{event_id}/participants/upload", response_model=UploadResponse)
async def upload_society_participants_excel(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    file: UploadFile = File(...),
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    content = await file.read()
    try:
        rows, errors = parse_participants_excel(content)
    except Exception as exc:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"Invalid Excel file. Please upload a valid .xlsx template. ({exc})",
        )

    created = 0
    skipped = 0
    for row in rows:
        email = row.pop("_email", "").strip().lower()
        raw_cert_type = row.pop("_cert_type", "participant")
        cert_type = _resolve_event_category_role(raw_cert_type, getattr(event, "category", None))

        existing = await Participant.find_one(
            Participant.event_id == event_id,
            Participant.email == email,
        )
        if existing:
            skipped += 1
            errors.append(f"Duplicate: {email}")
            continue

        p = Participant(
            event_id=event_id,
            club_id=society_id,
            email=email,
            registration_number=(
                row.get("Registration Number")
                or row.get("Reg No")
                or row.get("Roll No")
                or ""
            ),
            cert_type=cert_type,
            fields=row,
            source=ParticipantSource.EXCEL,
            verified=True,
        )
        await p.insert()
        created += 1

    current_count = await Participant.find(Participant.event_id == event_id).count()
    await event.set({"participant_count": current_count})

    return UploadResponse(created=created, skipped=skipped, errors=errors)


class ParticipantVerifyBody(BaseModel):
    verified: bool = True


@router.patch("/{society_id}/events/{event_id}/participants/{participant_id}/verify")
async def verify_society_participant(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    participant_id: PydanticObjectId,
    body: Optional[ParticipantVerifyBody] = None,
    _user: User = Depends(require_society_event_access),
):
    p = await Participant.get(participant_id)
    if not p or p.event_id != event_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Participant not found")
    new_verified = body.verified if body is not None else True
    await p.set({
        "verified": new_verified,
        "status": "accepted" if new_verified else "registered",
        "manually_verified": True,
    })
    return {"message": "Participant verification updated", "verified": new_verified}


@router.patch("/{society_id}/events/{event_id}/participants/{participant_id}/type")
async def update_society_participant_type(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    participant_id: PydanticObjectId,
    body: ParticipantTypeUpdate,
    _user: User = Depends(require_society_event_access),
):
    p = await Participant.get(participant_id)
    if not p or p.event_id != event_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Participant not found")
    await p.set({"cert_type": body.cert_type})
    return {"message": "Participant type updated"}


@router.delete("/{society_id}/events/{event_id}/participants/{participant_id}")
async def delete_society_participant(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    participant_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    p = await Participant.get(participant_id)
    if not p or p.event_id != event_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Participant not found")

    email = (p.email or "").strip().lower()
    if email:
        reg = await EventRegistration.find_one(
            EventRegistration.event_id == event_id,
            EventRegistration.student_email == email,
        )
        if reg:
            await reg.delete()

    event = await SocietyEvent.get(event_id)
    if event and event.participant_count > 0:
        await event.set({"participant_count": max(0, event.participant_count - 1)})

    certs = await Certificate.find(
        Certificate.event_id == event_id,
        Certificate.participant_id == str(p.id),
    ).to_list()
    for c in certs:
        if getattr(c, "status", None) != CertStatus.EMAILED.value:
            await c.delete()

    await p.delete()
    return {"message": "Participant deleted"}


@router.get("/{society_id}/events/{event_id}/participants/export")
async def export_society_event_participants(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_access),
):
    society = await ProfessionalSociety.get(society_id)
    if not society:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Professional Society not found")

    participants = await Participant.find(Participant.event_id == event_id).to_list()

    wb = Workbook()
    ws = wb.active
    ws.title = "Participants"

    headers = ["Name", "Email", "Reg No", "Type", "Source", "Status", "Added On"]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for p in participants:
        name = p.fields.get("Name", p.fields.get("name", "")) if p.fields else ""
        ws.append([
            name,
            p.email,
            p.registration_number or "",
            p.cert_type.replace("_", " ").title() if p.cert_type else "",
            p.source.value if hasattr(p.source, "value") else str(p.source),
            p.status or "",
            p.registered_at.strftime("%Y-%m-%d %H:%M:%S") if p.registered_at else "",
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"{society.slug}_event_{event_id}_participants.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ─── Attendance Endpoints ────────────────────────────────────────────────────

@router.post("/{society_id}/events/{event_id}/generate-qr", response_model=GenerateQRResponse)
@router.post("/{society_id}/events/{event_id}/attendance/qr", response_model=GenerateQRResponse)
async def generate_society_attendance_qr(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found.")

    if event.status == SocietyEventStatus.DRAFT:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Event is in Draft status. Please activate the event to generate attendance QR.",
        )
    if event.status == SocietyEventStatus.CLOSED:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Event is closed. Attendance QR cannot be generated for closed events.",
        )

    # QR generation is strictly allowed only on the event date
    if event.event_date:
        ist_offset = timedelta(hours=5, minutes=30)
        now_ist = (_datetime.now(timezone.utc) + ist_offset).date()
        event_dt = event.event_date
        event_ist = (event_dt.replace(tzinfo=timezone.utc) + ist_offset).date() if event_dt.tzinfo is None else (event_dt + ist_offset).date()
        if now_ist != event_ist and now_ist != event_dt.date():
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                "QR code generation is only allowed on the event day.",
            )

    current_count = getattr(event, "qr_generations_count", 0) or 0
    if current_count >= 3:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Maximum QR code generation limit (3 times) reached for this event.",
        )

    new_count = current_count + 1
    await event.set({"qr_generations_count": new_count})

    nonce = str(uuid.uuid4())
    issued_at = time.time()
    _store_nonce(nonce, issued_at)

    payload_dict = {
        "event_id": str(event_id),
        "club_id": str(society_id),
        "nonce": nonce,
        "issued_at": issued_at,
    }

    return GenerateQRResponse(
        payload=json.dumps(payload_dict),
        issued_at=issued_at,
        expires_in=QR_VALIDITY_SECONDS,
        qr_generations_count=new_count,
    )


@router.get("/{society_id}/events/{event_id}/attendance/active-session")
async def get_society_attendance_active_session(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found.")
    active_count = await Attendance.find(Attendance.event_id == event_id).count()
    return {"event_id": str(event_id), "attendance_count": active_count, "qr_generations_count": getattr(event, "qr_generations_count", 0)}


@router.post("/{society_id}/events/{event_id}/attendance/end-session")
async def end_society_attendance_session(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    return {"message": "Attendance session ended"}


@router.get("/{society_id}/events/{event_id}/attendance/records")
async def list_society_attendance_records(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    records = await Attendance.find(Attendance.event_id == event_id).to_list()
    return [
        {
            "id": str(r.id),
            "student_id": str(r.student_id),
            "student_name": r.student_name,
            "student_email": r.student_email,
            "registration_number": r.registration_number,
            "marked_at": r.marked_at,
            "feedback": r.feedback,
        }
        for r in records
    ]


# ─── Certificate Management (NO EVENT REPORT CHECK) ──────────────────────────

@router.post("/{society_id}/events/{event_id}/certificates/generate", response_model=GenerateResponse)
async def generate_society_certificates(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    background_tasks: BackgroundTasks,
    _user: User = Depends(require_society_event_access),
):
    event = await SocietyEvent.get(event_id)
    if not event or event.society_id != society_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Event not found")

    society = await ProfessionalSociety.get(society_id)
    if not society:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Professional Society not found")

    participants = await Participant.find(
        Participant.event_id == event_id,
        Participant.verified == True,
    ).to_list()

    existing_certs = await Certificate.find(Certificate.event_id == event_id).to_list()

    active_by_pid = {}
    retry_by_pid = {}
    for c in existing_certs:
        pid = c.participant_id
        if c.status in {CertStatus.PENDING, CertStatus.EMAILED}:
            active_by_pid[pid] = c
        elif c.status in {CertStatus.GENERATED, CertStatus.FAILED, CertStatus.REVOKED} and pid not in retry_by_pid:
            retry_by_pid[pid] = c

    queued = 0
    skipped_existing = 0
    skipped_no_template = 0
    year = _datetime.utcnow().year

    preset_docs = await RoleTemplatePreset.find(RoleTemplatePreset.is_active == True).to_list()
    preset_roles = {p.role_name for p in preset_docs}

    fp_docs = await FieldPosition.find(FieldPosition.event_id == event_id).to_list()
    fp_types = {fp.cert_type for fp in fp_docs if fp.template_filename}
    has_participant_fallback = "participant" in fp_types

    for p in participants:
        if p.id in active_by_pid:
            skipped_existing += 1
            continue

        raw_cert_type = p.cert_type or "participant"
        cert_type = _resolve_event_category_role(raw_cert_type, getattr(event, "category", None))
        normalized_role = cert_type.lower().replace(" ", "_").replace("-", "_")
        candidate_roles = [normalized_role]
        if normalized_role in {"volunteer", "student_volunteer"}:
            candidate_roles = ["student_volunteer", "volunteer"]

        has_preset = any(role in preset_roles for role in candidate_roles)
        has_fp = (cert_type in fp_types) or (raw_cert_type in fp_types) or (normalized_role in fp_types) or has_participant_fallback

        if not has_preset and not has_fp:
            skipped_no_template += 1
            continue

        existing_c = retry_by_pid.get(p.id)
        if existing_c:
            cert = existing_c
            cert.status = CertStatus.PENDING
            cert.snapshot.name = p.fields.get("Name") or cert.snapshot.name
            cert.snapshot.email = p.email
            cert.snapshot.registration_number = p.registration_number
            cert.snapshot.event_name = event.name
            cert.snapshot.club_name = society.name
            cert.snapshot.cert_type = cert_type
            await cert.save()
        else:
            cert_num = await generate_cert_number(society.slug or "society", year)
            cert = Certificate(
                cert_number=cert_num,
                participant_id=p.id,
                event_id=event_id,
                template_id=p.id,
                club_id=society_id,
                snapshot=CertSnapshot(
                    name=p.fields.get("Name") or "",
                    email=p.email,
                    registration_number=p.registration_number,
                    event_name=event.name,
                    club_name=society.name,
                    cert_type=cert_type,
                    issued_date=_datetime.utcnow(),
                ),
                status=CertStatus.PENDING,
            )
            await cert.insert()
            soc_cert = SocietyCertificate(
                cert_number=cert_num,
                participant_id=p.id,
                event_id=event_id,
                template_id=p.id,
                society_id=society_id,
                club_id=society_id,
                snapshot=cert.snapshot,
                status=CertStatus.PENDING,
            )
            await soc_cert.insert()

        background_tasks.add_task(_generate_one_limited, cert.id)
        queued += 1

    total = len(participants)
    msg = (
        f"Queued {queued} of {total} participant(s). "
        f"Skipped existing: {skipped_existing}. "
        f"Skipped missing template: {skipped_no_template}."
    )
    return GenerateResponse(
        queued_count=queued,
        queued=queued,
        skipped_existing=skipped_existing,
        skipped_no_template=skipped_no_template,
        total=total,
        message=msg,
    )


@router.get("/{society_id}/events/{event_id}/certificates", response_model=List[CertificateResponse])
async def list_society_certificates(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    status: Optional[str] = None,
    _user: User = Depends(require_society_event_access),
):
    query: Dict[str, Any] = {"event_id": event_id}
    if status:
        query["status"] = status
    certs = await Certificate.find(query).to_list()
    if not certs:
        soc_certs = await SocietyCertificate.find(query).to_list()
        if soc_certs:
            certs = soc_certs

    cert_ids = [c.id for c in certs]
    latest_error_by_cert: Dict[str, str] = {}
    if cert_ids:
        logs = await EmailLog.find(
            {"certificate_id": {"$in": cert_ids}}
        ).sort(-EmailLog.id).to_list()
        for log in logs:
            cid = str(log.certificate_id)
            if cid not in latest_error_by_cert and log.error_msg:
                latest_error_by_cert[cid] = log.error_msg

    return [_cert_resp(c, latest_error_by_cert) for c in certs]


@router.get("/{society_id}/events/{event_id}/certificates/counts")
@router.get("/{society_id}/events/{event_id}/certificates/check-all-status")
async def get_society_certificate_counts(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    certs = await Certificate.find(Certificate.event_id == event_id).to_list()
    if not certs:
        soc_certs = await SocietyCertificate.find(SocietyCertificate.event_id == event_id).to_list()
        if soc_certs:
            certs = soc_certs
    generated = sum(1 for c in certs if c.status in [CertStatus.GENERATED, CertStatus.EMAILED])
    emailed = sum(1 for c in certs if c.status == CertStatus.EMAILED)
    pending = sum(1 for c in certs if c.status == CertStatus.PENDING)
    failed = sum(1 for c in certs if c.status == CertStatus.FAILED)
    return {
        "total": len(certs),
        "generated": generated,
        "emailed": emailed,
        "pending": pending,
        "failed": failed,
    }


@router.post("/{society_id}/events/{event_id}/certificates/send-remaining")
@router.post("/{society_id}/events/{event_id}/certificates/email")
async def send_society_certificates_email(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    background_tasks: BackgroundTasks,
    _user: User = Depends(require_society_event_access),
):
    certs = await Certificate.find(
        Certificate.event_id == event_id,
        Certificate.status == CertStatus.GENERATED,
    ).to_list()

    for cert in certs:
        background_tasks.add_task(_send_email_limited, cert.id)

    return {"message": f"{len(certs)} email(s) queued for dispatch", "queued": len(certs)}


@router.get("/{society_id}/events/{event_id}/certificates/{cert_id}/download")
async def download_society_certificate(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    cert_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    cert = await Certificate.get(cert_id)
    if not cert or cert.event_id != event_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Certificate not found")
    if not cert.png_url:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Certificate PNG not generated yet")
    path = storage_url_to_path(cert.png_url)
    if not Path(path).exists():
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Certificate file not found on disk")
    return FileResponse(path, media_type="image/png", filename=f"{cert.cert_number}.png")


@router.post("/{society_id}/events/{event_id}/certificates/{cert_id}/resend")
@router.post("/{society_id}/events/{event_id}/certificates/{cert_id}/email")
async def resend_society_certificate_email(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    cert_id: PydanticObjectId,
    background_tasks: BackgroundTasks,
    _user: User = Depends(require_society_event_access),
):
    cert = await Certificate.get(cert_id)
    if not cert or cert.event_id != event_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Certificate not found")
    background_tasks.add_task(_send_email_limited, cert.id)
    return {"message": "Email re-queued"}


@router.post("/{society_id}/events/{event_id}/certificates/{cert_id}/generate")
async def regenerate_society_certificate(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    cert_id: PydanticObjectId,
    background_tasks: BackgroundTasks,
    _user: User = Depends(require_society_event_access),
):
    cert = await Certificate.get(cert_id)
    if not cert or cert.event_id != event_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Certificate not found")
    await cert.set({"status": CertStatus.PENDING})
    background_tasks.add_task(_generate_one_limited, cert.id)
    return {"message": "Certificate re-generation queued"}


@router.post("/{society_id}/events/{event_id}/certificates/download-all")
async def download_all_society_certificates(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    certs = await Certificate.find({
        "event_id": event_id,
        "status": {"$in": [CertStatus.GENERATED.value, CertStatus.EMAILED.value]},
    }).to_list()

    if not certs:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "No generated certificates available to download")

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for c in certs:
            if not c.png_url:
                continue
            path = storage_url_to_path(c.png_url)
            if Path(path).exists():
                arcname = f"{c.cert_number}_{c.snapshot.name or 'participant'}.png"
                zip_file.write(path, arcname=arcname)

    zip_buffer.seek(0)
    return StreamingResponse(
        zip_buffer,
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename=society_event_{event_id}_certificates.zip"},
    )


@router.get("/{society_id}/events/{event_id}/certificates/preview-layout")
@router.get("/{society_id}/events/{event_id}/certificates/preview-dynamic-layout")
async def preview_society_cert_layout(
    society_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_society_event_access),
):
    return {"message": "Preview layout endpoint ready"}


# ══════════════════════════════════════════════════════════════════════════════
# Society Coordinator Endpoints
# ══════════════════════════════════════════════════════════════════════════════

@coordinator_router.get("/stats")
async def coordinator_stats(
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    society = await ProfessionalSociety.get(current_user.society_id) if current_user.society_id else None
    total_events = await SocietyEvent.find(SocietyEvent.society_id == current_user.society_id).count() if current_user.society_id else 0
    active_members = await StudentSocietyMembership.find({
        "society_id": current_user.society_id,
        "status": SocietyMembershipStatus.APPROVED.value,
    }).count() if current_user.society_id else 0

    total_certs = await Certificate.find({"club_id": current_user.society_id}).count() if current_user.society_id else 0

    return {
        "society_name": society.name if society else "",
        "total_events": total_events,
        "active_members_count": active_members,
        "total_certificates": total_certs,
    }


@coordinator_router.get("/events")
async def coordinator_events(
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    if not current_user.society_id:
        return []
    events = await SocietyEvent.find(SocietyEvent.society_id == current_user.society_id).sort("-created_at").to_list()
    return [
        {
            "id": str(e.id),
            "name": e.name,
            "status": e.status.value if hasattr(e.status, "value") else str(e.status),
            "event_date": e.event_date,
        }
        for e in events
    ]


@coordinator_router.get("/memberships/requests")
async def get_membership_requests(
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    if not current_user.society_id and current_user.role != UserRole.SUPER_ADMIN:
        return []

    society_id = current_user.society_id
    query: dict[str, Any] = {}
    if society_id:
        query["society_id"] = society_id

    memberships = await StudentSocietyMembership.find(query).sort("-applied_at").to_list()
    return [
        {
            "id": str(m.id),
            "student_id": str(m.student_id),
            "student_name": m.student_name,
            "student_email": m.student_email,
            "society_id": str(m.society_id),
            "society_name": m.society_name,
            "status": m.status.value if hasattr(m.status, "value") else str(m.status),
            "office_bearer_role": m.office_bearer_role,
            "applied_at": m.applied_at,
            "review_note": m.review_note,
        }
        for m in memberships
    ]


@coordinator_router.put("/memberships/{membership_id}/status")
@coordinator_router.patch("/memberships/{membership_id}")
async def review_single_membership(
    membership_id: PydanticObjectId,
    body: ReviewMembershipRequest,
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    membership = await StudentSocietyMembership.get(membership_id)
    if not membership:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Membership request not found")

    if current_user.role == UserRole.PROFESSIONAL_SOCIETY_COORDINATOR:
        if current_user.society_id != membership.society_id:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "No access to this membership")

    action = (body.action or body.status or "").lower()
    if action in ["approve", "approved"]:
        membership.status = SocietyMembershipStatus.APPROVED
        act_text = "approved"
    elif action in ["reject", "rejected"]:
        membership.status = SocietyMembershipStatus.REJECTED
        act_text = "rejected"
    else:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Action must be 'approve' or 'reject'")

    membership.reviewed_by = str(current_user.id)
    membership.review_note = body.review_note
    membership.updated_at = _datetime.utcnow()
    await membership.save()

    return {"message": f"Membership {act_text} successfully"}


@coordinator_router.post("/memberships/bulk-approve")
@coordinator_router.patch("/memberships/requests/bulk-approve")
async def bulk_approve_memberships(
    body: BulkApproveMembershipsRequest,
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    society_id = current_user.society_id
    note = body.review_note.strip() if body.review_note else "Bulk approved by coordinator"
    now = _datetime.utcnow()
    reviewer = str(current_user.id)
    approved_count = 0

    for raw_id in body.membership_ids:
        try:
            m_id = ObjectId(raw_id)
        except Exception:
            continue

        membership = await StudentSocietyMembership.get(m_id)
        if not membership:
            continue

        if current_user.role == UserRole.PROFESSIONAL_SOCIETY_COORDINATOR and membership.society_id != society_id:
            continue

        membership.status = SocietyMembershipStatus.APPROVED
        membership.reviewed_by = reviewer
        membership.review_note = note
        membership.updated_at = now
        await membership.save()
        approved_count += 1

    return {
        "message": f"Successfully approved {approved_count} membership request(s)",
        "approved_count": approved_count,
    }


@coordinator_router.delete("/members/{identifier}")
async def remove_society_member(
    identifier: str,
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    if not current_user.society_id and current_user.role == UserRole.PROFESSIONAL_SOCIETY_COORDINATOR:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No society assigned to coordinator")

    society_id = current_user.society_id
    membership = None

    try:
        oid = ObjectId(identifier)
        membership = await StudentSocietyMembership.get(oid)
        if membership and current_user.role == UserRole.PROFESSIONAL_SOCIETY_COORDINATOR and membership.society_id != society_id:
            membership = None
    except Exception:
        pass

    if not membership and society_id:
        try:
            sid = ObjectId(identifier)
            membership = await StudentSocietyMembership.find_one({
                "society_id": society_id,
                "student_id": sid,
                "status": SocietyMembershipStatus.APPROVED.value,
            })
        except Exception:
            pass

    if not membership and society_id:
        membership = await StudentSocietyMembership.find_one({
            "society_id": society_id,
            "student_email": identifier.strip().lower(),
            "status": SocietyMembershipStatus.APPROVED.value,
        })

    if not membership:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Society membership not found")

    student_name = membership.student_name or "Student"
    society_name = membership.society_name or "the professional society"
    await membership.delete()

    return {"message": f"Successfully removed {student_name} from {society_name}"}


@coordinator_router.delete("/memberships/{membership_id}")
async def delete_society_membership(
    membership_id: PydanticObjectId,
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    membership = await StudentSocietyMembership.get(membership_id)
    if not membership:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Membership request not found")

    if current_user.role == UserRole.PROFESSIONAL_SOCIETY_COORDINATOR:
        if current_user.society_id != membership.society_id:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "No access to this membership")

    await membership.delete()
    return {"message": "Membership removed successfully"}


# ─── Office Bearers ──────────────────────────────────────────────────────────

@coordinator_router.get("/office-bearers")
async def get_society_office_bearers(
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    society_id = current_user.society_id
    if not society_id:
        return {"positions": DEFAULT_SOCIETY_OFFICE_BEARER_POSITIONS, "office_bearers": []}

    society = await ProfessionalSociety.get(society_id)
    positions = await _get_society_positions(society)

    members = await StudentSocietyMembership.find({
        "society_id": society_id,
        "office_bearer_role": {"$ne": None},
    }).to_list()

    return {
        "positions": positions,
        "office_bearers": [
            {
                "id": str(m.id),
                "student_id": str(m.student_id),
                "student_name": m.student_name,
                "student_email": m.student_email,
                "position": m.office_bearer_role,
            }
            for m in members
        ],
    }


@coordinator_router.post("/office-bearers/positions")
async def add_society_office_bearer_position(
    body: AddPositionBody,
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    if not current_user.society_id:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No society assigned to coordinator")

    society = await ProfessionalSociety.get(current_user.society_id)
    if not society:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Professional Society not found")

    new_position = (body.position or "").strip()
    if not new_position:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Position name cannot be empty")

    current_positions = await _get_society_positions(society)
    if any(p.lower() == new_position.lower() for p in current_positions):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Position '{new_position}' already exists")

    current_positions.append(new_position)
    society.office_bearer_positions = current_positions
    await society.save()

    return {
        "message": f"Position '{new_position}' added successfully",
        "positions": current_positions,
    }


@coordinator_router.delete("/office-bearers/positions/{position}")
async def delete_society_office_bearer_position(
    position: str,
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    if not current_user.society_id:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No society assigned to coordinator")

    society = await ProfessionalSociety.get(current_user.society_id)
    if not society:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Professional Society not found")

    pos_clean = position.strip()
    current_positions = await _get_society_positions(society)

    if pos_clean not in current_positions:
        raise HTTPException(status.HTTP_404_NOT_FOUND, f"Position '{pos_clean}' not found")

    holder = await StudentSocietyMembership.find_one({
        "society_id": current_user.society_id,
        "office_bearer_role": pos_clean,
    })
    if holder:
        holder.office_bearer_role = None
        holder.updated_at = _datetime.utcnow()
        await holder.save()

    current_positions.remove(pos_clean)
    society.office_bearer_positions = current_positions
    await society.save()

    return {
        "message": f"Position '{pos_clean}' removed",
        "positions": current_positions,
    }


@coordinator_router.post("/office-bearers/allocate")
@coordinator_router.post("/office-bearers")
async def allocate_society_office_bearer(
    body: AllocateOfficeBearerBody,
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    society_id = current_user.society_id
    if not society_id:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No society assigned to coordinator")

    society = await ProfessionalSociety.get(society_id)
    if not society:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Professional Society not found")

    position = (body.position or "").strip()
    valid_positions = await _get_society_positions(society)
    if position not in valid_positions:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"Invalid position. Must be one of: {', '.join(valid_positions)}",
        )

    query: dict[str, Any] = {
        "society_id": society_id,
        "status": SocietyMembershipStatus.APPROVED.value,
    }
    if body.student_id:
        try:
            query["student_id"] = ObjectId(body.student_id)
        except Exception:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Invalid student_id")
    elif body.student_email:
        query["student_email"] = body.student_email.strip().lower()
    else:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "student_id or student_email is required")

    membership = await StudentSocietyMembership.find_one(query)
    if not membership:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Student must be an approved member of this professional society to become an office bearer.",
        )

    # Exclusivity check across societies
    existing_ob = await StudentSocietyMembership.find_one({
        "student_id": membership.student_id,
        "office_bearer_role": {"$ne": None},
    })
    if existing_ob and (existing_ob.id != membership.id):
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"Student {membership.student_name or ''} is already an office bearer ({existing_ob.office_bearer_role}) in {existing_ob.society_name or 'another society'}.",
        )

    # Clear position from previous holder in this society
    previous_holder = await StudentSocietyMembership.find_one({
        "society_id": society_id,
        "office_bearer_role": position,
    })
    if previous_holder and previous_holder.id != membership.id:
        previous_holder.office_bearer_role = None
        await previous_holder.save()

    membership.office_bearer_role = position
    membership.updated_at = _datetime.utcnow()
    await membership.save()

    return {
        "message": f"Successfully allocated {position} to {membership.student_name}",
        "position": position,
        "student_id": str(membership.student_id),
    }


@coordinator_router.delete("/office-bearers/{position}")
async def remove_society_office_bearer(
    position: str,
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    society_id = current_user.society_id
    if not society_id:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No society assigned to coordinator")

    pos_clean = position.strip()
    holder = await StudentSocietyMembership.find_one({
        "society_id": society_id,
        "office_bearer_role": pos_clean,
    })
    if holder:
        holder.office_bearer_role = None
        holder.updated_at = _datetime.utcnow()
        await holder.save()

    return {"message": f"Office bearer position '{pos_clean}' unassigned"}


# ─── Active Members Tab Endpoint (Cumulative Participation) ──────────────────

@coordinator_router.get("/active-members")
async def get_society_active_members(
    current_user: User = Depends(require_role(UserRole.PROFESSIONAL_SOCIETY_COORDINATOR, UserRole.SUPER_ADMIN)),
):
    if not current_user.society_id:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No society assigned to coordinator")

    society_id = current_user.society_id

    # 1. All events for this society
    all_events = await SocietyEvent.find(SocietyEvent.society_id == society_id).to_list()
    total_events = len(all_events)
    event_ids = [ev.id for ev in all_events]

    # 2. Approved memberships
    approved_memberships = await StudentSocietyMembership.find({
        "society_id": society_id,
        "status": SocietyMembershipStatus.APPROVED.value,
    }).to_list()

    approved_emails = set()
    for m in approved_memberships:
        if m.student_email:
            approved_emails.add(m.student_email.strip().lower())

    # 3. Participants
    participants = await Participant.find({
        "event_id": {"$in": event_ids},
    }).to_list() if event_ids else []

    email_to_events: dict[str, set] = {}
    for p in participants:
        email_key = (p.email or "").strip().lower()
        if email_key:
            email_to_events.setdefault(email_key, set()).add(str(p.event_id))

    all_participant_emails = set(email_to_events.keys())

    # 4. User documents
    user_query_clauses = []
    if approved_emails:
        user_query_clauses.append({"email": {"$in": list(approved_emails)}})
    if all_participant_emails:
        user_query_clauses.append({"email": {"$in": list(all_participant_emails)}})

    users = await User.find({
        "role": UserRole.STUDENT,
        "$or": user_query_clauses,
    }).to_list() if user_query_clauses else []

    member_user_ids = [m.student_id for m in approved_memberships if getattr(m, "student_id", None)]
    member_users = await User.find({"_id": {"$in": member_user_ids}}).to_list() if member_user_ids else []

    user_map_by_email: dict[str, User] = {}
    user_map_by_id: dict[str, User] = {}
    for u in users:
        email_key = (u.email or "").strip().lower()
        if email_key:
            user_map_by_email[email_key] = u
        user_map_by_id[str(u.id)] = u
    for u in member_users:
        user_map_by_id[str(u.id)] = u
        email_key = (u.email or "").strip().lower()
        if email_key:
            user_map_by_email[email_key] = u

    seen_emails: set[str] = set()
    result = []

    for m in approved_memberships:
        email_key = (m.student_email or "").strip().lower()
        if not email_key or email_key in seen_emails:
            continue
        seen_emails.add(email_key)

        u = user_map_by_email.get(email_key) or user_map_by_id.get(str(m.student_id))
        ev_count = len(email_to_events.get(email_key, set()))
        ratio = f"{ev_count}/{total_events}" if total_events > 0 else f"{ev_count}/0"
        pct = round((ev_count / total_events) * 100) if total_events > 0 else 0

        result.append({
            "membership_id": str(m.id),
            "student_id": str(m.student_id) if m.student_id else None,
            "name": m.student_name or (u.name if u else None) or "—",
            "registration_number": getattr(m, "registration_number", None) or (u.registration_number if u else None) or "—",
            "department": getattr(m, "department", None) or (u.department if u else None) or "—",
            "email": m.student_email,
            "membership_status": "Yes",
            "office_bearer_role": m.office_bearer_role,
            "participation_ratio": ratio,
            "participation_percentage": pct,
            "events_participated": ev_count,
        })

    for p_email in all_participant_emails:
        if p_email in seen_emails:
            continue
        seen_emails.add(p_email)
        u = user_map_by_email.get(p_email)
        ev_count = len(email_to_events.get(p_email, set()))
        ratio = f"{ev_count}/{total_events}" if total_events > 0 else f"{ev_count}/0"
        pct = round((ev_count / total_events) * 100) if total_events > 0 else 0

        result.append({
            "membership_id": None,
            "student_id": str(u.id) if u else None,
            "name": u.name if u else p_email,
            "registration_number": u.registration_number if u else "—",
            "department": u.department if u else "—",
            "email": p_email,
            "membership_status": "No",
            "office_bearer_role": None,
            "participation_ratio": ratio,
            "participation_percentage": pct,
            "events_participated": ev_count,
        })

    return {
        "total_events": total_events,
        "total_approved_members": len(approved_memberships),
        "total_students": len(result),
        "members": result,
    }
