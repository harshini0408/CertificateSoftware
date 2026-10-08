import sys
path = r'C:\D\CertificateSoftware\backend\app\domains\club\routers.py'
with open(path, 'r', encoding='utf-8') as f:
    text = f.read()

old_endpoint = '''@router.get("/{club_id}/events/{event_id}/participants/export")
async def export_event_participants(
    club_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_club_access)
):
    import io
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from ...models.participant import Participant

    club = await Club.get(club_id)
    if not club:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Club not found")
    
    participants = await Participant.find({"event_id": str(event_id)}).to_list()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Participants"
    
    headers = ["Name", "Email", "Reg No", "Type", "Source", "Status", "Added On"]
    ws.append(headers)
    
    for p in participants:
        ws.append([
            p.student_name,
            p.student_email,
            p.registration_number or "",
            p.participant_type,
            p.source,
            p.status,
            p.created_at.strftime("%Y-%m-%d %H:%M:%S") if p.created_at else ""
        ])
        
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=participants_{event_id}.xlsx"}
    )'''

new_endpoint = '''@router.get("/{club_id}/events/{event_id}/participants/export")
async def export_event_participants(
    club_id: PydanticObjectId,
    event_id: PydanticObjectId,
    _user: User = Depends(require_club_access)
):
    import io
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from ...models.participant import Participant

    club = await Club.get(club_id)
    if not club:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Club not found")

    participants = await Participant.find(
        Participant.event_id == event_id
    ).to_list()

    wb = Workbook()
    ws = wb.active
    ws.title = "Participants"

    headers = ["Name", "Email", "Reg No", "Type", "Source", "Status", "Added On"]
    ws.append(headers)

    # Bold header row
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for p in participants:
        name = p.fields.get("Name", p.fields.get("name", ""))
        ws.append([
            name,
            p.email,
            p.registration_number or "",
            p.cert_type.replace("_", " ").title() if p.cert_type else "",
            p.source.value if hasattr(p.source, "value") else str(p.source),
            p.status or "",
            p.registered_at.strftime("%Y-%m-%d %H:%M:%S") if p.registered_at else ""
        ])

    # Auto-fit columns
    for col in ws.columns:
        max_len = max((len(str(cell.value)) for cell in col if cell.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=participants_{event_id}.xlsx"}
    )'''

text = text.replace(old_endpoint, new_endpoint)
with open(path, 'w', encoding='utf-8') as f:
    f.write(text)
print("Done.")
